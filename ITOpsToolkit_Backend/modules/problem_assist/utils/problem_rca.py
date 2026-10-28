import os
import sys
import json
import asyncio
import logging
import datetime
from datetime import datetime, timedelta
from typing import List

from pydantic import BaseModel, Field

from llm_guard.vault import Vault as LLMGuardVault
from llm_guard import scan_prompt, scan_output
from llm_guard.output_scanners.bias import MatchType as OutBiasMatchType
from llm_guard.input_scanners.toxicity import MatchType as InpToxMatchType
from llm_guard.output_scanners.toxicity import MatchType as OutToxMatchType
from llm_guard.input_scanners.gibberish import MatchType as InpGibMatchType
from llm_guard.output_scanners.gibberish import MatchType as OutGibMatchType
from llm_guard.input_scanners.prompt_injection import MatchType as InpPIMatchType
from llm_guard.output_scanners import Toxicity as OutToxicity, Gibberish as OutGibberish, MaliciousURLs, Deanonymize, Bias
from llm_guard.input_scanners import Toxicity as InpToxicity, Gibberish as InpGibberish, BanCode as InpBanCode, PromptInjection, Anonymize, Secrets

SCRIPT_PATH = os.path.dirname(__file__)
ROOT_PATH = SCRIPT_PATH.split("modules")[0]
DB_PATH = os.path.join(ROOT_PATH, "database")
AUTH_PATH = os.path.join(ROOT_PATH, "modules", "auth", "utils")
REACTIVE_PATH = os.path.join(SCRIPT_PATH, "reactive_problem")
PA_CONFIG_PATH = os.path.join(os.path.dirname(SCRIPT_PATH), "config")

sys.path.append(ROOT_PATH)
sys.path.append(DB_PATH)
sys.path.append(SCRIPT_PATH)
sys.path.append(AUTH_PATH)
sys.path.append(REACTIVE_PATH)

from task import ServiceNowTasks
from problems import ServiceNowProblems
from incidents import ServiceNowIncidents
from attachments import ServiceNowAttachments

from knowledge_assistant import KnowledgeAssistant

import core_crud
import db_session
import system_config_pa as scf

from llm import LLM
from usage_tracker import track_llm_tokens
from pro_logger import log_event, application_logger

from docx import Document


TS_FORMAT = '%Y-%m-%d %H:%M:%S'


# ── Response schemas ──────────────────────────────────────────────────────────

class AskLLMResponse(BaseModel):
    """Generic LLM agent response."""

    response: str = Field(default="", description="The response to the user's query.")


class ProblemAnalysisResponse(BaseModel):
    """Single-call structured response covering all problem analysis fields."""

    # corrective_action: str = Field(
    #     default="N/A",
    #     description=()
    # )
    # preventive_action: str = Field(
    #     default="N/A",
    #     description=(
    #         "Top 3 preventive actions to avoid recurrence. "
    #         "Include 'N/A' if no preventive actions are available."
    #     )
    # )
    corrective_action: str = Field(
    default="N/A",
    description=(
        "Generate a detailed corrective action report. "
        "Begin with an introductory paragraph. "
        "Then provide numbered corrective actions with detailed explanations. "
        "Include recovery activities, troubleshooting actions, validation steps, restoration actions, repair actions, tools used, commands used if available, and outcomes achieved. "
        "Your response should not be less than 5 to 7 sentences, with proper bullet points and sub-bullet points, and headings if required."
        )
    )

    preventive_action: str = Field(
        default="N/A",
        description=(
            "Generate a detailed preventive action report. "
            "Begin with an introductory paragraph. "
            "Then provide exactly 3 numbered preventive actions. "
            "Each preventive action must include a title and a detailed explanation of implementation approach, expected benefits, risk reduction, monitoring considerations, and recurrence prevention strategy. "
            "Your response should not be less than 5 to 7 sentences, with proper bullet points and sub-bullet points, and headings if required."
        )
    )
    root_cause_summary: str = Field(
        default="N/A",
        description=(
            "Generate a detailed Root Cause Analysis Summary. "
            "Provide a comprehensive overview of the entire scenario. "
            "Structure the response using the following sections:\n"
            "Problem Overview\n"
            "Root Causes\n"
            "Resolution Steps\n"
            "Impact\n"
            "Explain the sequence of events, contributing factors, technical findings, operational findings, corrective actions taken, and resulting business impact. "
            "The response should be at least 200 words when sufficient information is available. "
            "Your response should not be less than 5 to 7 sentences, with proper bullet points and sub-bullet points, and headings if required."
            "Use a professional RCA report style."
        )
    )
    corrective_action_summary: str = Field(
        default="",
        description="Summary of the corrective action in two sentences."
    )
    preventive_action_summary: str = Field(
        default="",
        description="Summary of the preventive action in two sentences."
    )
    impact: str = Field(
        default="",
        description="Specific impacts of the problem in two to three sentences."
    )
    cause: str = Field(
        default="",
        description="Cause of the problem summarized in 2-3 lines."
    )
    monitoring_miss_flag: bool = Field(
        default=False,
        description="It should be 'true' if caused by a monitoring miss, else 'false'."
    )
    manual_mistake: bool = Field(
        default=False,
        description="It should be 'true' if caused by a manual mistake, else 'false'."
    )
    process_issue: bool = Field(
        default=False,
        description="It should be 'true' if caused by a process issue, else 'false'."
    )
    technical_rca: str = Field(
        default="N/A",
        description=(
                        "Perform a detailed Technical 5-Why Analysis. "
                        "Focus only on technical causes such as application defects, infrastructure failures, "
                        "configuration issues, database problems, network failures, monitoring tool issues, "
                        "automation failures, software bugs, hardware failures, integrations, and capacity limitations. "
                        "Do not include process, governance, communication, training, ownership, or human factors. "
                        "Provide exactly 5 levels in the format below:\n\n"
                        "Question 1: <Question>\n"
                        "Answer 1: <Answer>\n"
                        "Question 2: <Question>\n"
                        "Answer 2: <Answer>\n"
                        "Question 3: <Question>\n"
                        "Answer 3: <Answer>\n"
                        "Question 4: <Question>\n"
                        "Answer 4: <Answer>\n"
                        "Question 5: <Question>\n"
                        "Answer 5: <Answer>\n"
                        "Each answer must logically lead to the next question.\n"
                        "Use actual line breaks and do not return escaped newline characters (\\n)."
    )
    )
    process_rca: str = Field(
        default="N/A",
        description=(
                        "Perform a detailed Process 5-Why Analysis. "
                        "Focus only on process, governance, operational, communication, training, "
                        "change management, documentation, approval workflow, monitoring strategy, "
                        "ownership, compliance, escalation, and organizational factors. "
                        "Do not include technical failures, software defects, infrastructure issues, "
                        "database problems, network failures, or configuration errors. "
                        "Provide exactly 5 levels in the format below:\n\n"
                        "Question 1: <Question>\n"
                        "Answer 1: <Answer>\n"
                        "Question 2: <Question>\n"
                        "Answer 2: <Answer>\n"
                        "Question 3: <Question>\n"
                        "Answer 3: <Answer>\n"
                        "Question 4: <Question>\n"
                        "Answer 4: <Answer>\n"
                        "Question 5: <Question>\n"
                        "Answer 5: <Answer>\n"
                        "Each answer must logically lead to the next question.\n"
                        "Use actual line breaks and do not return escaped newline characters (\\n)."
        )
    )


class PreventiveTask(BaseModel):
    """A single preventive action task."""

    short_description: str = Field(default="", description="Title of the preventive action.")
    description: str = Field(default="", description="Detailed explanation of the preventive action.")
    cmdb_ci: str = Field(default="", description="Configuration item (CI) associated with the action.")
    due_date: int = Field(default=0, description="Time required to complete the action in minutes from now.")
    work_notes: str = Field(default="", description="Known error details if any, else empty string.")
    assignment_group: str = Field(default="", description="Most relevant assignment group from the provided list.")
    justification: str = Field(default="", description="Justification for the due_date value.")


class PreventiveTasksResponse(BaseModel):
    """Structured list of preventive action tasks."""

    tasks: List[PreventiveTask] = Field(default_factory=list, description="List of preventive action tasks.")

class IncidentSummaryResponse(BaseModel):
    """Grouped incident summary returned as plain text."""

    summary: str = Field(
        default="",
        description=(
            "Group incidents based on similar issues and provide one paragraph "
            "for each group. Each paragraph should include the issue, cause, "
            "and resolution. If only one incident exists, summarize that incident. "
            "Return a single consolidated text summary."
        )
    )


class AttachmentSummary(BaseModel):
    """Summary of a single attachment."""

    summary: str = Field(
        default="",
        description="Concise summary of the attachment content."
    )


class AttachmentSummaryResponse(BaseModel):
    """Collection of attachment summaries."""

    attachments: List[AttachmentSummary] = Field(
        default_factory=list,
        description="List of attachment summaries."
    )
       
# ── Agent factories ───────────────────────────────────────────────────────────

def ask_llm_agent(sys_prompts_config):
    configs = scf.get_system_config()
    ala_sys_prompt = sys_prompts_config.get("ask_llm_agent", None)
    (description, instruction) = scf.apply_system_prompt_guardrail(ala_sys_prompt)
    llm_tag = ala_sys_prompt.get("llm_tag", configs["llm_tag"]) if ala_sys_prompt else configs["llm_tag"]

    llm_obj = LLM()
    llm_obj.set_litellm_environment(llm_tag)
    client = llm_obj.client

    if not client:
        log_str = "ERROR: Model not configured for AskLLMAgent."
        log_event(application_logger,
                  logging.ERROR,
                  log_str,
                  module_name = "problem_assist",
                  component_type = "ask_llm_agent",
                  component_name = "rca_api")
        return (None, None)

    client._default_options["temperature"] = 0.0
    client._litellm_params["response_format"] = AskLLMResponse

    return client.create_agent(
        name = "AskLLMAgent",
        instructions = instruction,
        description = description,
    ), client

def problem_analysis_agent(sys_prompts_config):
    configs = scf.get_system_config()
    ala_sys_prompt = sys_prompts_config.get("ask_llm_agent", None)
    (description, instruction) = scf.apply_system_prompt_guardrail(ala_sys_prompt)
    llm_tag = ala_sys_prompt.get("llm_tag", configs["llm_tag"]) if ala_sys_prompt else configs["llm_tag"]

    llm_obj = LLM()
    llm_obj.set_litellm_environment(llm_tag)
    client = llm_obj.client

    if not client:
        log_str = "ERROR: Model not configured for ProblemAnalysisAgent."
        log_event(application_logger,
                  logging.ERROR,
                  log_str,
                  module_name = "problem_assist",
                  component_type = "problem_analysis_agent",
                  component_name = "rca_api")
        return (None, None)

    client._default_options["temperature"] = 0.0
    client._litellm_params["response_format"] = ProblemAnalysisResponse

    return client.create_agent(
        name = "ProblemAnalysisAgent",
        instructions = instruction,
        description = description,
    ), client

def preventive_tasks_agent(sys_prompts_config):
    configs = scf.get_system_config()
    ala_sys_prompt = sys_prompts_config.get("ask_llm_agent", None)
    (description, instruction) = scf.apply_system_prompt_guardrail(ala_sys_prompt)
    llm_tag = ala_sys_prompt.get("llm_tag", configs["llm_tag"]) if ala_sys_prompt else configs["llm_tag"]

    llm_obj = LLM()
    llm_obj.set_litellm_environment(llm_tag)
    client = llm_obj.client

    if not client:
        log_event(application_logger,
                  logging.ERROR,
                  "ERROR: Model not configured for PreventiveTasksAgent.",
                  module_name="problem_assist",
                  component_type="preventive_tasks_agent",
                  component_name="rca_api")
        return (None, None)

    client._default_options["temperature"] = 0.0
    client._litellm_params["response_format"] = PreventiveTasksResponse

    return client.create_agent(
        name="PreventiveTasksAgent",
        instructions=instruction,
        description=description,
    ), client

def incident_summary_agent(sys_prompts_config):
    configs = scf.get_system_config()

    ala_sys_prompt = sys_prompts_config.get("ask_llm_agent", None)

    (description, instruction) = scf.apply_system_prompt_guardrail(
        ala_sys_prompt
    )

    llm_tag = (
        ala_sys_prompt.get("llm_tag", configs["llm_tag"])
        if ala_sys_prompt
        else configs["llm_tag"]
    )

    llm_obj = LLM()
    llm_obj.set_litellm_environment(llm_tag)

    client = llm_obj.client

    if not client:
        return (None, None)

    client._default_options["temperature"] = 0.0
    client._litellm_params["response_format"] = IncidentSummaryResponse

    return client.create_agent(
        name="IncidentSummaryAgent",
        instructions=instruction,
        description=description,
    ), client

def attachment_summary_agent(sys_prompts_config):
    configs = scf.get_system_config()

    ala_sys_prompt = sys_prompts_config.get("ask_llm_agent", None)

    (description, instruction) = scf.apply_system_prompt_guardrail(
        ala_sys_prompt
    )

    llm_tag = (
        ala_sys_prompt.get("llm_tag", configs["llm_tag"])
        if ala_sys_prompt
        else configs["llm_tag"]
    )

    llm_obj = LLM()
    llm_obj.set_litellm_environment(llm_tag)

    client = llm_obj.client

    if not client:
        return (None, None)

    client._default_options["temperature"] = 0.0
    client._litellm_params["response_format"] = AttachmentSummaryResponse

    return client.create_agent(
        name="AttachmentSummaryAgent",
        instructions=instruction,
        description=description,
    ), client
    
class RCAGenerator:
    def __init__(self):
        configs = scf.get_system_config()
        itsm_tags = configs.get("itsm_tags", [])

        db_local_session = db_session.get_db()
        itsm_details = core_crud.get_itsm_connector_details_with_creds(db_local_session, itsm_tags[0])
        db_local_session.close()

        if itsm_details is None:
            log_event(application_logger,
                       logging.ERROR,
                       f"Failed to retrieve ITSM connector details for tag '{itsm_tags[0]}'. "
                       "Check vault connectivity and secret configuration.",
                       module_name="problem_assist",
                       component_type="rca_generator",
                       component_name="itsm_connector")
        self.itsm_url = itsm_details.itsm_url
        self.itsm_tag = itsm_tags[0]

        assist_backstory_f = open(PA_CONFIG_PATH + "/task_backstory.txt", "r")
        self.assistant_backstory = assist_backstory_f.read()
        assist_backstory_f.close()


    def create_problem_prompt(self, problem, attachment_summary, incident_summary):
        problem_ticket_prompt_background = 'There are the core problem details - '
        problem_ticket_prompt_background += '\n* Short description - ' + problem['short_description']
        problem_ticket_prompt_background += '\n* Description - ' + problem['description']
        problem_ticket_prompt_background += '\n* Comments and work notes - ' + problem['comments_and_work_notes']
        problem_ticket_prompt_background += '\n* These are the details of problem\'s attachment' + '\n' + attachment_summary
        problem_ticket_prompt_background += '\n* Here is the detailed overview of incidents related to the Problem' + '\n' \
                                            + incident_summary

        problem_ticket_prompt = 'You are a Technical SME who works on the problem ticket management.\
                                Below are the details of a particular problem\n ' '--------------------------------\n\n' \
                                + problem_ticket_prompt_background + '\n--------------------------------\n\n'

        return problem_ticket_prompt

    def create_work_notes(self, problem, attachment_summary, incident_summary):
        work_notes_background = 'These are the problem details - '
        work_notes_background += '\n* Short description - ' + problem['short_description']
        work_notes_background += '\n* Description - ' + problem['description']
        work_notes_background += '\n* Comments and work notes - ' + problem['comments_and_work_notes']
        work_notes_background += '\n* These are the details of problem\'s attachment' + '\n' + attachment_summary
        work_notes_background += '\n* Here is the detailed overview of incidents related to the Problem' + '\n' \
                                            + incident_summary

        work_notes = 'Gen-AI leverage the below information to generate the root cause analysis\n ' \
                    '--------------------------------\n\n' \
                    + work_notes_background + \
                    '\n--------------------------------\n\n'

        return work_notes


    def get_problem_details(self, number = None, start_ts = None, end_ts = None):
        problem_obj = ServiceNowProblems()
        problem_response = problem_obj.get_problems_details(number = number, start_ts = None, end_ts = None)

        return problem_response


    def get_problem_incident_details(self, problem_number):
        incident_obj = ServiceNowIncidents()
        incident_response = incident_obj.get_incidents_details(problem_number)
        inc_data = {}
        inc_numbers = ''
        inc_created_on_list = []
        inc_resolved_on_list = []
        inc_priority_list = []

        if incident_response == []:
            return []

        for inc in incident_response:
            inc_atch_obj = ServiceNowAttachments()

            if inc_numbers == '':
                inc_numbers = inc['number']
            else:
                inc_numbers += ',' + inc['number']

            inc_created_on_list.append(inc['sys_created_on'])
            inc_resolved_on_list.append(inc['resolved_at'])
            inc_priority_list.append(inc['priority'])

            inc_atch_mtdt = inc_atch_obj.get_attachments_metadata('incident', inc['sys_id'])
            inc_file_paths = inc_atch_obj.download_attachments(inc_atch_mtdt)

            inc['inc_attachments_content'] = []

            for path in inc_file_paths:
                inc_atch_content = inc_atch_obj.read_attachment(path)
                if inc_atch_content is not None:
                    inc['inc_attachments_content'].append(str(inc_atch_content))
                else:
                    continue

        inc_created_on = min(inc_created_on_list)
        inc_resolved_on = max(inc_resolved_on_list)
        inc_priority = min(inc_priority_list)

        inc_data['inc_details'] = incident_response
        inc_data['inc_numbers'] = inc_numbers
        inc_data['inc_created_on'] = inc_created_on
        inc_data['inc_resolved_on'] = inc_resolved_on
        inc_data['inc_priority'] = inc_priority

        return inc_data

    def get_tasks_details(self, problem_sys_id):
        task_obj = ServiceNowTasks()
        task_response = task_obj.get_task_details(problem_sys_id)

        return task_response

    def get_attachments_content(self, problem_sys_id):
        atch_obj = ServiceNowAttachments()
        atch_mtdt = atch_obj.get_attachments_metadata('problem', problem_sys_id)
        file_paths = atch_obj.download_attachments(atch_mtdt)

        content = []

        for path in file_paths:
            atch_content = atch_obj.read_attachment(path)
            if atch_content is not None:
                content.append(str(atch_content))
            else:
                continue

        return content

    async def get_incident_summary_structured(self, prompt):
        configs = scf.get_system_config()

        sys_prompts_config = configs["sys_prompts_config"]

        IncidentSummaryAgent, agent_client = incident_summary_agent(
            sys_prompts_config
        )

        if IncidentSummaryAgent is None:
            return ""

        llm_guards = configs["llm_guards"]

        vault = LLMGuardVault()

        input_llm_guard_dict = {
            "toxicity": InpToxicity(
                threshold=0.5,
                match_type=InpToxMatchType.SENTENCE
            ),
            "gibberish": InpGibberish(
                match_type=InpGibMatchType.FULL
            ),
            "bancode": InpBanCode(),
            "prompt_injection": PromptInjection(
                threshold=0.5,
                match_type=InpPIMatchType.FULL
            ),
            "anonymization": Anonymize(vault),
        }

        input_scanners_list = []

        for item in llm_guards:
            if item in input_llm_guard_dict:
                input_scanners_list.append(
                    input_llm_guard_dict[item]
                )

        sanitized_prompt, _, _ = scan_prompt(
            input_scanners_list,
            prompt
        )

        agent_client.usage = {}

        outcome = await IncidentSummaryAgent.run(
            sanitized_prompt
        )

        track_llm_tokens(
            usage=agent_client.usage,
            agent="IncidentSummaryAgent",
            llm_tag=sys_prompts_config.get(
                "ask_llm_agent",
                {}
            ).get(
                "llm_tag",
                configs["llm_tag"]
            ),
            module="problem_assist",
            user_id=self.user_id,
        )

        if isinstance(outcome, IncidentSummaryResponse):
            return outcome.summary

        try:
            outcome_dict = json.loads(str(outcome))
            return outcome_dict.get("summary", "")
        except Exception:
            return str(outcome)
        
    async def get_attachment_summary_structured(self, attachments):
        """
        Summarize attachments using AttachmentSummaryAgent.

        Returns:
            str: Concatenated attachment summaries.
        """

        if not attachments:
            return ""

        configs = scf.get_system_config()
        sys_prompts_config = configs["sys_prompts_config"]

        AttachmentSummaryAgent, agent_client = attachment_summary_agent(
            sys_prompts_config
        )

        if AttachmentSummaryAgent is None:
            log_event(
                application_logger,
                logging.ERROR,
                "AttachmentSummaryAgent initialization failed.",
                module_name="problem_assist",
                component_type="get_attachment_summary_structured",
                component_name="rca_api",
            )
            return ""

        prompt = (
            "Analyze each attachment independently.\n"
            "Provide a concise but meaningful summary for each attachment.\n"
            "Do not combine attachment summaries.\n"
            "Return one summary object per attachment in the same order.\n\n"
        )

        for idx, attachment in enumerate(attachments, start=1):
            prompt += (
                f"Attachment {idx}:\n"
                f"{attachment}\n\n"
            )

        try:
            #
            # LLM Guard Input Scanning
            #
            llm_guards = configs["llm_guards"]

            vault = LLMGuardVault()

            input_llm_guard_dict = {
                "toxicity": InpToxicity(
                    threshold=0.5,
                    match_type=InpToxMatchType.SENTENCE,
                ),
                "gibberish": InpGibberish(
                    match_type=InpGibMatchType.FULL,
                ),
                "bancode": InpBanCode(),
                "prompt_injection": PromptInjection(
                    threshold=0.5,
                    match_type=InpPIMatchType.FULL,
                ),
                "anonymization": Anonymize(vault),
            }

            input_scanners_list = []

            for item in llm_guards:
                if item in input_llm_guard_dict:
                    input_scanners_list.append(
                        input_llm_guard_dict[item]
                    )

            sanitized_prompt, _, _ = scan_prompt(
                input_scanners_list,
                prompt,
            )

            #
            # LLM Call
            #
            agent_client.usage = {}

            outcome = await AttachmentSummaryAgent.run(
                sanitized_prompt
            )

            #
            # Token Tracking
            #
            track_llm_tokens(
                usage=agent_client.usage,
                agent="AttachmentSummaryAgent",
                llm_tag=sys_prompts_config.get(
                    "ask_llm_agent",
                    {}
                ).get(
                    "llm_tag",
                    configs["llm_tag"],
                ),
                module="problem_assist",
                user_id=self.user_id,
            )

            #
            # Structured response handling
            #
            if isinstance(outcome, AttachmentSummaryResponse):

                result = []

                for idx, attachment in enumerate(
                    outcome.attachments,
                    start=1,
                ):
                    result.append(
                        f"Attachment {idx}:\n{attachment.summary}"
                    )

                return "\n\n".join(result)

            #
            # Fallback parsing
            #
            try:

                outcome_dict = json.loads(str(outcome))

                result = []

                for idx, attachment in enumerate(
                    outcome_dict.get("attachments", []),
                    start=1,
                ):
                    result.append(
                        f"Attachment {idx}:\n"
                        f"{attachment.get('summary', '')}"
                    )

                return "\n\n".join(result)

            except Exception:
                return str(outcome)

        except Exception as e:

            log_event(
                application_logger,
                logging.ERROR,
                f"Attachment summary generation failed: {str(e)}",
                module_name="problem_assist",
                component_type="get_attachment_summary_structured",
                component_name="rca_api",
            )

            return ""
    
    async def get_incident_summary(self, incidents):
        """
        Generates a grouped incident summary in plain text.

        Similar incidents are grouped together and summarized into
        issue, cause, and resolution paragraphs.
        """

        if not incidents:
            return ""

        inc_summary_prompt_background = (
            "\nHere are the details for related incidents -"
        )

        for inc in incidents:

            inc_summary_prompt_background += (
                "\n\n--------------------------------\n\n"
            )

            inc_summary_prompt_background += (
                f"* Number -\n{inc['number']}\n\n"
            )

            inc_summary_prompt_background += (
                f"* Short Description -\n{inc['short_description']}\n\n"
            )

            inc_summary_prompt_background += (
                f"* Description -\n{inc['description']}\n\n"
            )

            inc_summary_prompt_background += (
                f"* Closure Notes -\n{inc['close_notes']}\n\n"
            )

            #
            # Use structured attachment summary agent
            #
            if len(inc.get("inc_attachments_content", [])) != 0:

                attachment_summary = await self.get_attachment_summary(
                    inc["inc_attachments_content"]
                )

                inc_summary_prompt_background += (
                    "* Appendix -\n"
                    "There are several attachments present in the incident.\n"
                    "Here is the summary of the attachments:\n"
                )

                inc_summary_prompt_background += (
                    attachment_summary + "\n"
                )

            inc_summary_prompt_background += (
                "\n--------------------------------\n"
            )

        #
        # Final summarization instructions
        #
        inc_summary_prompt_background += (
            "\n\nPerform the following tasks on the above data:\n"
            "1. Group incidents based on similar issues using the short description and description.\n"
            "2. For each group, provide a concise summary covering:\n"
            "   - Issue\n"
            "   - Cause\n"
            "   - Resolution\n"
            "3. Use exactly one paragraph per incident group.\n"
            "4. If only one incident exists, provide a summary for that incident.\n"
            "5. Return a single consolidated summary string.\n"
        )

        try:
            return await self.get_incident_summary_structured(
                inc_summary_prompt_background
            )
        except Exception as exc:

            log_event(application_logger,
                      logging.ERROR,
                      f"Failed to generate incident summary: {str(exc)}",
                      module_name="problem_assist",
                      component_type="get_incident_summary",
                      component_name="rca_api"
                      )

            return ""
        
    async def get_attachment_summary(self, attachments):
        return await self.get_attachment_summary_structured(attachments)



    async def obtain_all_analysis(self, problem_prompt):
        """Single LLM call replacing 10 separate obtain_* calls."""
        prompt = (
            problem_prompt +
            "\n\nBased on the problem details above, provide a thorough analysis covering ALL of the following:\n"
            "1. corrective_action: Generate a detailed corrective action report. "
                "The output must contain an introductory paragraph followed by numbered corrective actions. "
                "For each corrective action explain what was done, why it was done, how it was performed, any tools used, validation steps, recovery actions, restoration activities, repair activities, and outcomes achieved. "
                "Do not provide a one-line answer or brief summary.\n"

            "2. preventive_action: Generate a detailed preventive action report. "
                "Start with an introductory paragraph. "
                "Then provide exactly 3 numbered preventive actions. "
                "For each preventive action include a title and a detailed explanation describing implementation approach, risk reduction, expected benefits, monitoring considerations, and recurrence prevention strategy. "
                "Do not provide short bullet points or one-line recommendations. "
                "If no preventive actions can be identified return N/A.\n"
            "3. root_cause_summary: Generate a detailed Root Cause Analysis Summary. "
                "Provide a comprehensive overview of the entire scenario. "
                "Structure the response using the following sections:\n"
                "Problem Overview:\n"
                "Root Causes:\n"
                "Resolution Steps:\n"
                "Impact:\n"
                "Explain the sequence of events, contributing factors, technical and operational findings, corrective actions taken, and resulting impact. "
                "Do not provide a short summary or 2-3 line response. "
                "Use a professional RCA report style with detailed explanations."
            "4. corrective_action_summary: Summarize the corrective action in two sentences.\n"
            "5. preventive_action_summary: Summarize the preventive action in two sentences.\n"
            "6. impact: List the specific impacts of the problem in two to three sentences "
            "   (do not include generic impacts).\n"
            "7. cause: Summarize the cause of the problem in 2-3 lines.\n"
            "8. monitoring_miss_flag: Return true if the issue was caused by a monitoring miss or lack of effective monitoring. Otherwise return false.\n IMPORTANT: Return actual boolean values only. monitoring_miss_flag must be true or false, manual_mistake must be true or false, and process_issue must be true or false. Do not return strings such as 'true', 'false', 'yes', 'no', 'boolean_0', or 'boolean_1'.\n"
            "9. manual_mistake: Return true if the issue was caused by human error, manual intervention, incorrect execution, or operator mistake. Otherwise return false.\n IMPORTANT: Return actual boolean values only. monitoring_miss_flag must be true or false, manual_mistake must be true or false, and process_issue must be true or false. Do not return strings such as 'true', 'false', 'yes', 'no', 'boolean_0', or 'boolean_1'.\n"
            "10. process_issue: Return true if the issue was caused by a process gap, procedural weakness, governance issue, missing SOP, inadequate change management, poor documentation, or operational process failure. Otherwise return false.\n IMPORTANT: Return actual boolean values only. monitoring_miss_flag must be true or false, manual_mistake must be true or false, and process_issue must be true or false. Do not return strings such as 'true', 'false', 'yes', 'no', 'boolean_0', or 'boolean_1'.\n"
            "11. technical_rca: Perform a detailed Technical 5 Why analysis. Focus only on technical causes such as application defects, infrastructure failures, database issues, configuration issues, networking failures, software bugs, automation failures, hardware failures, integrations, and capacity limitations. Do not include governance, process, training, communication, ownership, or operational factors.\n"
            "12. process_rca: Perform a detailed Process 5 Why analysis. Focus only on process gaps, governance failures, SOP deficiencies, documentation issues, change management problems, communication gaps, training deficiencies, operational weaknesses, monitoring strategy gaps, and ownership issues. Do not include software defects, infrastructure failures, network issues, database problems, or configuration errors.\n"
            "\n"
            "Technical RCA and Process RCA must not reuse the same root causes.\n"
        )

        configs = scf.get_system_config()
        sys_prompts_config = configs["sys_prompts_config"]

        if sys_prompts_config is None:
            log_event(application_logger,
                      logging.ERROR,
                      "Guardrail note - System prompts not defined. Cannot proceed.",
                      module_name="problem_assist",
                      component_type="obtain_all_analysis",
                      component_name="rca_api")
            return ProblemAnalysisResponse()

        llm_guards = configs["llm_guards"]

        ProblemAnalysisAgent, agent_client = problem_analysis_agent(sys_prompts_config)

        if ProblemAnalysisAgent is None:
            log_event(application_logger,
                      logging.ERROR,
                      "ERROR: Model not configured for ProblemAnalysisAgent.",
                      module_name="problem_assist",
                      component_type="obtain_all_analysis",
                      component_name="rca_api")
            return ProblemAnalysisResponse()

        vault = LLMGuardVault()
        input_llm_guard_dict = {
            "toxicity": InpToxicity(threshold=0.5, match_type=InpToxMatchType.SENTENCE),
            "gibberish": InpGibberish(match_type=InpGibMatchType.FULL),
            "bancode": InpBanCode(),
            "prompt_injection": PromptInjection(threshold=0.5, match_type=InpPIMatchType.FULL),
            "anonymization": Anonymize(vault),
        }

        input_scanners_list = []
        for item in llm_guards:
            if item in input_llm_guard_dict:
                input_scanners_list.append(input_llm_guard_dict[item])

        sanitized_prompt, _, _ = scan_prompt(input_scanners_list, prompt)

        agent_client.usage = {}
        outcome = await ProblemAnalysisAgent.run(sanitized_prompt)

        track_llm_tokens(
            usage=agent_client.usage,
            agent="ProblemAnalysisAgent",
            llm_tag=sys_prompts_config.get("ask_llm_agent", {}).get("llm_tag", configs["llm_tag"]),
            module="problem_assist",
            user_id=self.user_id,
        )

        if isinstance(outcome, ProblemAnalysisResponse):
            return outcome

        # Fallback: try to parse from string
        try:
            outcome_dict = json.loads(str(outcome))
            return ProblemAnalysisResponse(**outcome_dict)
        except Exception:
            return ProblemAnalysisResponse()

    async def obtain_preventive_act_json(self, preventive_action, known_error, a_grp) -> List[PreventiveTask]:
        """Single structured LLM call returning a list of PreventiveTask objects."""
        prompt = 'Below are the Preventive actions\n'
        prompt += preventive_action
        prompt += '\nBelow is the list of Assignment groups\n'
        prompt += a_grp

        if 'No relevant KEDB articles found in repository' not in known_error:
            prompt += '\nBelow are the Known errors\n'
            prompt += known_error

        prompt += '\n\n' + self.assistant_backstory

        configs = scf.get_system_config()
        sys_prompts_config = configs["sys_prompts_config"]

        if sys_prompts_config is None:
            log_event(application_logger,
                      logging.ERROR,
                      "Guardrail note - System prompts not defined. Cannot proceed.",
                      module_name="problem_assist",
                      component_type="obtain_preventive_act_json",
                      component_name="rca_api")
            return []

        llm_guards = configs["llm_guards"]

        PreventiveAgent, agent_client = preventive_tasks_agent(sys_prompts_config)

        if PreventiveAgent is None:
            log_event(application_logger,
                      logging.ERROR,
                      "ERROR: Model not configured for PreventiveTasksAgent.",
                      module_name="problem_assist",
                      component_type="obtain_preventive_act_json",
                      component_name="rca_api")
            return []

        vault = LLMGuardVault()
        input_llm_guard_dict = {
            "toxicity": InpToxicity(threshold=0.5, match_type=InpToxMatchType.SENTENCE),
            "gibberish": InpGibberish(match_type=InpGibMatchType.FULL),
            "bancode": InpBanCode(),
            "prompt_injection": PromptInjection(threshold=0.5, match_type=InpPIMatchType.FULL),
            "anonymization": Anonymize(vault),
        }

        input_scanners_list = []
        for item in llm_guards:
            if item in input_llm_guard_dict:
                input_scanners_list.append(input_llm_guard_dict[item])

        sanitized_prompt, _, _ = scan_prompt(input_scanners_list, prompt)

        agent_client.usage = {}
        outcome = await PreventiveAgent.run(sanitized_prompt)

        track_llm_tokens(
            usage=agent_client.usage,
            agent="PreventiveTasksAgent",
            llm_tag=sys_prompts_config.get("ask_llm_agent", {}).get("llm_tag", configs["llm_tag"]),
            module="problem_assist",
            user_id=self.user_id,
        )

        if isinstance(outcome, PreventiveTasksResponse):
            return outcome.tasks

        # Fallback: try to parse from string
        try:
            outcome_dict = json.loads(str(outcome))
            return PreventiveTasksResponse(**outcome_dict).tasks
        except Exception:
            return []

    async def ask_llm(self, query):
        """Run AskLLMAgent with LLM Guard input/output scanning."""
        configs = scf.get_system_config()
        sys_prompts_config = configs["sys_prompts_config"]

        if sys_prompts_config is None:
            log_str = "Guardrail note - System prompts not defined for any agent. Cannot proceed."
            log_event(application_logger,
                      logging.ERROR,
                      log_str,
                      module_name = "problem_assist",
                      component_type = "ask_llm",
                      component_name = "rca_api")
            return log_str

        llm_guards = configs["llm_guards"]

        AskLLMAgent, ask_client = ask_llm_agent(sys_prompts_config = sys_prompts_config)

        if AskLLMAgent is None:
            log_str = "ERROR: Model not configured for AskLLMAgent."
            log_event(application_logger,
                      logging.ERROR,
                      log_str,
                      module_name = "problem_assist",
                      component_type = "ask_llm",
                      component_name = "rca_api")
            return log_str

        vault = LLMGuardVault()
        input_llm_guard_dict = {
            "toxicity": InpToxicity(threshold = 0.5, match_type = InpToxMatchType.SENTENCE),
            "gibberish": InpGibberish(match_type = InpGibMatchType.FULL),
            "bancode": InpBanCode(),
            "prompt_injection": PromptInjection(threshold = 0.5, match_type = InpPIMatchType.FULL),
            "anonymization": Anonymize(vault),
        }

        output_llm_guard_dict = {
            "toxicity": OutToxicity(threshold = 0.5, match_type = OutToxMatchType.SENTENCE),
            "malicious_urls": MaliciousURLs(threshold = 0.7),
            "gibberish": OutGibberish(match_type = OutGibMatchType.FULL),
            "anonymization": Deanonymize(vault),
            "bias": Bias(threshold = 0.5, match_type = OutBiasMatchType.FULL),
        }

        input_scanners_list = []
        output_scanners_list = []

        for item in llm_guards:
            if item in input_llm_guard_dict:
                input_scanners_list.append(input_llm_guard_dict[item])
            if item in output_llm_guard_dict:
                output_scanners_list.append(output_llm_guard_dict[item])

        sanitized_prompt, _, _ = scan_prompt(input_scanners_list, query)

        ask_client.usage = {}

        outcome = await AskLLMAgent.run(sanitized_prompt)

        track_llm_tokens(
            usage = ask_client.usage,
            agent = "AskLLMAgent",
            llm_tag = sys_prompts_config.get("ask_llm_agent", {}).get("llm_tag", configs["llm_tag"]),
            module = "problem_assist",
            user_id = self.user_id,
        )

        if hasattr(outcome, 'response'):
            outcome_str = outcome.response
        else:
            try:
                outcome_str = json.loads(str(outcome)).get("response", str(outcome))
            except (json.JSONDecodeError, TypeError):
                outcome_str = str(outcome)

        outcome_str, _, _ = scan_output(output_scanners_list, sanitized_prompt, outcome_str)
        return outcome_str

    def update_rca_document(self, data):
        mapping_file = PA_CONFIG_PATH + "/mapping.json"
        mapping_dict = json.load(open(mapping_file, 'r'))

        filename = 'Root_Cause_Analysis.docx'
        source = PA_CONFIG_PATH + '/' + filename
        destination = REACTIVE_PATH + '/output/'

        os.makedirs(destination, exist_ok=True)
        os.system(f"cp {source} {destination}")
        file_destination = destination + filename
        os.system(f"chmod 666 {file_destination}")
        doc = Document(file_destination)

        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                for k, v in mapping_dict.items():
                    run.text = run.text.replace(k, data[v])

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for k, v in mapping_dict.items():
                        cell.text = cell.text.replace(k, data[v])

        doc.save(file_destination)
        return file_destination


    async def main(self, number, user_id):
        self.user_id = user_id
        problem_list = self.get_problem_details(number = number)
        if not problem_list:
            return (False, f"Problem {number} not found in ITSM")
        problem = problem_list[0]

        incident_response = self.get_problem_incident_details(problem['number'])
        if not incident_response:
            log_event(application_logger,
                      logging.WARNING,
                      f"No incidents found for problem {problem['number']}. Skipping.",
                      module_name = "problem_assist",
                      component_type = "main",
                      component_name = "rca_api")
            return (False, f"No incidents found for problem {number}")

        problem['related_incidents'] = incident_response['inc_details']

        incident_list = incident_response['inc_numbers']
        incident_created_on = incident_response['inc_created_on']
        incident_resolved_on = incident_response['inc_resolved_on']
        incident_priority = incident_response['inc_priority']

        problem['related_tasks'] = self.get_tasks_details(problem['sys_id'])
        problem['attachments_content'] = self.get_attachments_content(problem['sys_id'])
        
        incident_summary = await self.get_incident_summary(problem['related_incidents'])
        attachment_summary = await self.get_attachment_summary(problem['related_incidents'][0]['inc_attachments_content'])

        problem_prompt = self.create_problem_prompt(problem, attachment_summary, incident_summary)

        # Single LLM call for all analysis fields
        analysis = await self.obtain_all_analysis(problem_prompt)

        corrective_action      = analysis.corrective_action
        preventive_action      = analysis.preventive_action
        root_cause_analysis_summary = analysis.root_cause_summary
        corrective_action_summary   = analysis.corrective_action_summary
        preventive_action_summary   = analysis.preventive_action_summary
        impact_of_the_issue    = analysis.impact
        cause_of_the_issue     = analysis.cause
        monitor_miss_flag      = analysis.monitoring_miss_flag
        manual_mistake         = analysis.manual_mistake
        process_issue          = analysis.process_issue

        problem_obj = ServiceNowProblems()
        work_notes = self.create_work_notes(problem, attachment_summary, incident_summary) 

        data = {}
        data['corrective_action'] = corrective_action
        data['preventive_action'] = preventive_action
        data['root_cause_summary'] = root_cause_analysis_summary
        data['corrective_action_summary'] = corrective_action_summary
        data['preventive_action_summary'] = preventive_action_summary
        data['impact'] = impact_of_the_issue
        data['cause'] = cause_of_the_issue
        data['incident_numbers'] = incident_list
        data['incident_created_on'] = incident_created_on
        data['incident_resolved_on'] = incident_resolved_on
        data['company'] = ''
        data['description'] = problem['description']
        data['short_description'] = problem['short_description'] + '\n'
        data['number'] = problem['number']
        data['current_utc_date_time'] = (datetime.now()).strftime(TS_FORMAT)
        data['highest_incident_priority'] = incident_priority
        data['technical_rca'] = 'N/A'
        data['process_rca'] = 'N/A'

        data['monitoring_miss_flag'] = 'Yes' if monitor_miss_flag else 'No'
        data['manual_mistake'] = 'Yes' if manual_mistake else 'No'
        data['process_issue'] = 'Yes' if process_issue else 'No'

        # Technical RCA
        if monitor_miss_flag or manual_mistake:
            data['technical_rca'] = analysis.technical_rca

        # Process RCA
        if manual_mistake or process_issue:
            data['process_rca'] = analysis.process_rca

        # Neither flag set
        if not monitor_miss_flag and not manual_mistake and not process_issue:
            data['technical_rca'] = analysis.technical_rca
            
        try:
            assignment_group = problem['assignment_group']['display_value']
        except:
            assignment_group = ''

        try:
            assigned_to = problem['assigned_to']['display_value']
        except:
            assigned_to = ''

        if preventive_action.strip().upper() == "N/A":
            data['known_error'] = "\nNo relevant known errors found."

        else:
            quest = 'These are the problem details\n' + problem['short_description'] + '\n' + problem['description']
            quest += '\n to which below preventive actions are suggested to avoid this particular issue in future.\n'
            quest += preventive_action
            quest += "\n Use only the knowledge article identifiers explicitly provided in the source data. \
                Do not generate, infer, abbreviate, rename, or invent article references. \
                Every reference must exactly match a KB number that exists in the provided knowledge article content. \
                If no KB reference is available, do not create one and state that No relevant known errors found."

            ka_obj = KnowledgeAssistant()
            known_error = ka_obj.get_compressed_response(quest)

            if 'No relevant KEDB articles found in repository' in known_error:
                data['known_error'] = '\nNo relevant known errors found.'
            else:
                data['known_error'] = str(known_error)

        task_obj = ServiceNowTasks()

        if preventive_action.strip().upper() == "N/A":
            pass
        else:
            a_grp = task_obj.get_assignment_grp()
            ptask_list = await self.obtain_preventive_act_json(preventive_action, known_error, a_grp)

            for ptask in ptask_list:
                c_time = datetime.now()
                d_time = c_time + timedelta(minutes=ptask.due_date)
                due_date_str = d_time.strftime(TS_FORMAT)

                task_obj.create_task(
                    short_description=ptask.short_description,
                    number=problem['number'],
                    description=ptask.description,
                    assignment_group=ptask.assignment_group,
                    due_date=due_date_str,
                    cmdb_ci=ptask.cmdb_ci,
                    work_notes=ptask.work_notes,
                )

        file_destination = self.update_rca_document(data)

        attachment_obj = ServiceNowAttachments()
        attach = attachment_obj.attach_attachment('Root_Cause_Analysis.docx', problem['sys_id'])

        update_worknotes = problem_obj.update_problem(problem['sys_id'], work_notes, table = 'problem')

        new_task = task_obj.create_task(number = problem['number'], assignment_group = assignment_group, assigned_to = assigned_to)

        os.system(f"rm {file_destination}")

        return (True, f"RCA created successfully for {number}")

if __name__ == '__main__':
    rca_generator = RCAGenerator()
    asyncio.run(rca_generator.main(number = "PRB0041198", user_id = "system"))
